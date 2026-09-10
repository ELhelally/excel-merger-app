import streamlit as st
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import pandas as pd
import os
import glob
import zipfile
import datetime
import io

# إعدادات الصفحة
st.set_page_config(page_title="دمج كشوف الحضور والانصراف", page_icon="📊", layout="wide")

st.title("📊 أداة دمج وتنسيق كشوف الحضور والانصراف (التنسيق المعتمد)")
st.write("قم برفع ملف الـ ZIP للحصول على التنسيق المعتمد بالكامل مع الترتيب التلقائي بأرقام العمال.")

uploaded_zip = st.file_uploader("اختر ملف الـ ZIP (مثل: employees.zip)", type=["zip"])
output_custom_name = st.text_input("📝 اكتب اسم ملف الإكسل الناتج:", value="كشف_حضور_وانصراف_شهر_مارس_المجمع")

DAY_MAP = {
    'السبت': 'Sat', 'الأحد': 'Sun', 'الاحد': 'Sun',
    'الإثنين': 'Mon', 'الاثنين': 'Mon', 'الانتين': 'Mon', 'الاتنين': 'Mon',
    'الثلاثاء': 'Tue', 'الأربعاء': 'Wed', 'الاربعاء': 'Wed',
    'الخميس': 'Thu', 'الجمعة': 'Fri', 'الجمعه': 'Fri'
}

def parse_time_24(val, is_end_time=False):
    if pd.isna(val) or val is None:
        return None
    val_str = str(val).strip()
    if not val_str or val_str == '-':
        return 0
    
    has_pm = 'م' in val_str
    has_am = 'ص' in val_str
    clean_num = ''.join([c for c in val_str if c.isdigit() or c == ':'])
    
    if not clean_num:
        return val_str
    
    try:
        hrs = int(clean_num.split(':')[0])
        if has_pm and hrs < 12:
            hrs += 12
        elif is_end_time and hrs < 12 and not has_am and hrs <= 11:
            hrs += 12
        return hrs
    except:
        return val_str

def parse_date_custom(val):
    if pd.isna(val) or val is None or str(val).strip() in ['', '-', 'None']:
        return None
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%d-%b-%y")
    
    val_str = str(val).split(' ')[0].strip()
    try:
        dt = pd.to_datetime(val_str, dayfirst=True, errors='coerce')
        if pd.notnull(dt):
            return dt.strftime("%d-%b-%y")
    except:
        pass
    return val_str

def clean_dash_to_zero(val):
    if pd.isna(val) or val is None:
        return 0
    val_str = str(val).strip()
    if val_str in ['-', '—', '']:
        return 0
    try:
        if '.' in val_str:
            return float(val_str)
        return int(val_str)
    except:
        return val_str

def process_employee_sheet(file_path):
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        all_rows = list(sheet.iter_rows(values_only=True))
        
        emp_name, emp_id, emp_role = None, None, None
        header_row_idx = None
        
        for r_idx, row in enumerate(all_rows):
            if any('اليوم' in str(c) for c in row if c) and any('التاريخ' in str(c) for c in row if c):
                header_row_idx = r_idx
                break
                
            for cell in row:
                if cell is None: continue
                cell_str = str(cell).strip()
                
                if 'الاسم' in cell_str and not emp_name:
                    emp_name = cell_str.split(':')[-1].strip() if ':' in cell_str else None
                    if not emp_name:
                        idx = list(row).index(cell)
                        if idx + 1 < len(row): emp_name = str(row[idx+1]).strip()

                if any(k in cell_str for k in ['رقم العامل', 'رقم الموظف', 'الكود']) and not emp_id:
                    emp_id = cell_str.split(':')[-1].strip() if ':' in cell_str else None
                    if not emp_id:
                        idx = list(row).index(cell)
                        if idx + 1 < len(row): emp_id = str(row[idx+1]).strip()

        fname = os.path.basename(file_path)
        if '-' in fname:
            f_parts = fname.replace('.xlsx', '').replace('.xls', '').split('-')
            if not emp_name and len(f_parts) >= 1: emp_name = f_parts[0].strip()
            if not emp_id and len(f_parts) >= 2: emp_id = f_parts[1].strip()

        # تحويل رقم العامل لرقم صحيح للترتيب
        try:
            emp_id_clean = int(re.sub(r'\D', '', str(emp_id))) if emp_id else 999999
        except:
            emp_id_clean = 999999

        rows_data = []
        if header_row_idx is not None:
            for r in all_rows[header_row_idx + 1:]:
                if not any(r): continue
                
                day_raw = str(r[0]).strip() if r[0] is not None else ""
                
                if any(term in day_raw for term in ['إجمالي', 'اجمالي', 'ملخص', 'عدد أيام', 'التقدير العام']):
                    if 'إجمالي' in day_raw or 'اجمالي' in day_raw: break
                    continue
                    
                if r[0] is None and r[1] is None: continue

                day_en = DAY_MAP.get(day_raw, day_raw)
                date_formatted = parse_date_custom(r[1])
                
                time_from = parse_time_24(r[6] if len(r) > 6 else None, is_end_time=False)
                time_to = parse_time_24(r[7] if len(r) > 7 else None, is_end_time=True)
                
                trans_val = clean_dash_to_zero(r[8] if len(r) > 8 else None)
                food_val = clean_dash_to_zero(r[9] if len(r) > 9 else None)
                
                process_num = r[3] if len(r) > 3 and str(r[3]).strip() not in ['-', '—'] else None
                process_name = r[4] if len(r) > 4 and str(r[4]).strip() not in ['-', '—'] else None
                location = r[5] if len(r) > 5 and str(r[5]).strip() not in ['-', '—'] else None
                supervisor = r[10] if len(r) > 10 else None

                is_work_day = 1 if (process_name or process_num or (time_from and str(time_from) != '0')) else None
                if 'إجازة' in str(process_name) or 'اجازة' in str(process_name) or day_en == 'Fri':
                    is_work_day = None

                rows_data.append({
                    'اليوم': day_en,
                    'التاريخ': date_formatted,
                    'رقم العامل': emp_id_clean if emp_id_clean != 999999 else emp_id,
                    'رقم العملية': process_num,
                    'العملية': process_name,
                    'المكان': location,
                    'من': time_from if is_work_day else None,
                    'إلى': time_to if is_work_day else None,
                    'الإنتقالات': trans_val if is_work_day else None,
                    'بدل غذاء': food_val if is_work_day else None,
                    'توقيع المشرف': supervisor if is_work_day else None,
                    'أيام العمل': is_work_day
                })
        return pd.DataFrame(rows_data)
    except Exception as e:
        st.error(f"خطأ في الملف {file_path}: {e}")
        return pd.DataFrame()

if uploaded_zip is not None:
    if st.button("🚀 دمج الملفات وتنسيق الجدول بالكامل"):
        with st.spinner("جاري التجهيز وترتيب أرقام العمال وتطبيق التنسيق..."):
            extract_dir = "./temp_extracted"
            os.makedirs(extract_dir, exist_ok=True)
            
            with zipfile.ZipFile(uploaded_zip, 'r') as zip_ref:
                zip_ref.extractall(extract_dir)
                
            all_files = glob.glob(f"{extract_dir}/**/*.xlsx", recursive=True) + glob.glob(f"{extract_dir}/**/*.xls", recursive=True)
            
            dfs = [process_employee_sheet(f) for f in all_files]
            master_df = pd.concat(dfs, ignore_index=True)
            
            # ترتيب الموظفين تصاعدياً برقم العامل (رقم الموظف)
            if 'رقم العامل' in master_df.columns:
                master_df['emp_sort_key'] = pd.to_numeric(master_df['رقم العامل'], errors='coerce')
                master_df = master_df.sort_values(by=['emp_sort_key'], ascending=True, kind='stable').drop(columns=['emp_sort_key'])

            cols_order = ['اليوم', 'التاريخ', 'رقم العامل', 'رقم العملية', 'العملية', 'المكان', 'من', 'إلى', 'الإنتقالات', 'بدل غذاء', 'توقيع المشرف', 'أيام العمل']
            for col in cols_order:
                if col not in master_df.columns:
                    master_df[col] = None
            master_df = master_df[cols_order]
            
            output_path = "temp_output.xlsx"
            master_df.to_excel(output_path, index=False)
            
            wb = openpyxl.load_workbook(output_path)
            ws = wb.active
            ws.sheet_view.rightToLeft = True
            
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            friday_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
            
            thin_border = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
            align_center = Alignment(horizontal='center', vertical='center')
            
            for col_num in range(1, len(cols_order) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_center

            for row_num in range(2, ws.max_row + 1):
                day_val = str(ws.cell(row=row_num, column=1).value or '')
                is_friday = (day_val == 'Fri' or day_val == 'الجمعة' or day_val == 'الجمعه')
                
                for col_num in range(1, len(cols_order) + 1):
                    cell = ws.cell(row=row_num, column=col_num)
                    cell.alignment = align_center
                    cell.border = thin_border
                    
                    if is_friday:
                        cell.fill = friday_fill

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

            output_buffer = io.BytesIO()
            wb.save(output_buffer)
            output_buffer.seek(0)
            
            clean_filename = output_custom_name.strip()
            if not clean_filename.endswith(".xlsx"):
                clean_filename += ".xlsx"
            
            st.success(f"✅ تم الدمج والترتيب التلقائي برقم العامل بنجاح!")
            
            st.download_button(
                label=f"📥 تحميل الملف المرتب المنسق: {clean_filename}",
                data=output_buffer,
                file_name=clean_filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
